# services/excel_export_service.py

import os
from datetime import datetime
from typing import List, Dict, Optional
from PyQt6.QtCore import QDate
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


class ExcelExportService:
    """Сервис для экспорта переработок в Excel"""

    def __init__(self):
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        self.header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        self.total_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
        self.subtotal_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    def export_overtimes(self, overtimes: List[Dict], start_date: QDate, end_date: QDate,
                         file_path: str, department: Optional[str] = None,
                         division: Optional[str] = None) -> str:
        """
        Экспортирует переработки в Excel файл

        Args:
            overtimes: список переработок
            start_date: начальная дата периода
            end_date: конечная дата периода
            file_path: путь для сохранения файла
            department: отдел (для фильтрации)
            division: подразделение (для фильтрации)

        Returns:
            str: путь к сохраненному файлу
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Переработки"

        # Заголовок отчета
        title = f"Отчет по переработкам за период {start_date.toString('dd.MM.yyyy')} - {end_date.toString('dd.MM.yyyy')}"
        if department:
            title += f" (Отдел: {department})"
        if division:
            title += f" (Подразделение: {division})"

        ws.merge_cells('A1:G1')
        title_cell = ws['A1']
        title_cell.value = title
        title_cell.font = Font(size=14, bold=True)
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        title_cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        title_cell.border = self.thin_border

        # Заголовки колонок
        headers = ['ФИО', 'Описание', 'Дата', 'Начало', 'Конец', 'Продолжительность (ч)', 'Примечание']
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col_num)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = self.thin_border

        # Группируем переработки по сотрудникам
        employees_data = {}
        for ot in overtimes:
            emp_name = ot['user']
            if emp_name not in employees_data:
                employees_data[emp_name] = {
                    'overtimes': [],
                    'total_hours': 0.0
                }

            # Преобразуем продолжительность из строки в число
            duration_str = ot['duration'].replace(',', '.')
            try:
                hours = float(duration_str)
            except:
                hours = 0.0

            employees_data[emp_name]['overtimes'].append(ot)
            employees_data[emp_name]['total_hours'] += hours

        # Сортируем сотрудников по имени
        sorted_employees = sorted(employees_data.items(), key=lambda x: x[0])

        current_row = 3
        grand_total = 0.0

        for emp_name, emp_data in sorted_employees:
            # Сортируем переработки сотрудника по дате
            emp_overtimes = sorted(emp_data['overtimes'],
                                   key=lambda x: QDate.fromString(x['date'], "dd.MM.yyyy").toPyDate())

            for ot in emp_overtimes:
                # ФИО
                ws.cell(row=current_row, column=1).value = emp_name

                # Описание
                ws.cell(row=current_row, column=2).value = ot['description']

                # Дата
                ws.cell(row=current_row, column=3).value = ot['date']

                # Начало
                ws.cell(row=current_row, column=4).value = ot['start_time']

                # Конец
                ws.cell(row=current_row, column=5).value = ot['end_time']

                # Продолжительность
                duration_cell = ws.cell(row=current_row, column=6)
                duration_cell.value = float(ot['duration'].replace(',', '.'))
                duration_cell.number_format = '0.0'

                # Примечание (проект/задача)
                note = []
                if ot.get('project'):
                    note.append(f"Проект: {ot['project']}")
                if ot.get('task'):
                    note.append(f"Задача: {ot['task']}")
                ws.cell(row=current_row, column=7).value = '; '.join(note) if note else ''

                # Применяем стили
                for col in range(1, 8):
                    cell = ws.cell(row=current_row, column=col)
                    cell.alignment = Alignment(horizontal='left', vertical='center')
                    cell.border = self.thin_border

                current_row += 1

            # Добавляем строку с итогом по сотруднику
            if emp_overtimes:
                # Объединяем ячейки для ФИО
                ws.merge_cells(f'A{current_row}:B{current_row}')
                total_label = ws.cell(row=current_row, column=1)
                total_label.value = f"ИТОГО по сотруднику {emp_name}:"
                total_label.font = Font(bold=True)
                total_label.fill = self.subtotal_fill
                total_label.alignment = Alignment(horizontal='right', vertical='center')
                total_label.border = self.thin_border

                # Пустые ячейки
                for col in [3, 4, 5]:
                    cell = ws.cell(row=current_row, column=col)
                    cell.border = self.thin_border
                    cell.fill = self.subtotal_fill

                # Итоговое количество часов
                total_hours_cell = ws.cell(row=current_row, column=6)
                total_hours_cell.value = emp_data['total_hours']
                total_hours_cell.font = Font(bold=True)
                total_hours_cell.fill = self.subtotal_fill
                total_hours_cell.alignment = Alignment(horizontal='center', vertical='center')
                total_hours_cell.border = self.thin_border
                total_hours_cell.number_format = '0.0'

                # Примечание
                note_cell = ws.cell(row=current_row, column=7)
                note_cell.border = self.thin_border
                note_cell.fill = self.subtotal_fill

                grand_total += emp_data['total_hours']
                current_row += 1

                # Пустая строка между сотрудниками
                current_row += 1

        # Общий итог
        if grand_total > 0:
            current_row += 1
            ws.merge_cells(f'A{current_row}:E{current_row}')
            total_label = ws.cell(row=current_row, column=1)
            total_label.value = "ОБЩИЙ ИТОГ ПО ВСЕМ СОТРУДНИКАМ:"
            total_label.font = Font(bold=True, size=12)
            total_label.fill = self.total_fill
            total_label.alignment = Alignment(horizontal='right', vertical='center')
            total_label.border = self.thin_border

            total_hours_cell = ws.cell(row=current_row, column=6)
            total_hours_cell.value = grand_total
            total_hours_cell.font = Font(bold=True, size=12)
            total_hours_cell.fill = self.total_fill
            total_hours_cell.alignment = Alignment(horizontal='center', vertical='center')
            total_hours_cell.border = self.thin_border
            total_hours_cell.number_format = '0.0'

        # Автоматическая ширина колонок
        for col in range(1, 8):
            max_length = 0
            column = get_column_letter(col)
            for row in range(1, current_row + 1):
                cell = ws[f'{column}{row}']
                if cell.value:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

        # Сохраняем файл
        wb.save(file_path)
        return file_path

    def generate_filename(self, start_date: QDate, end_date: QDate,
                          department: Optional[str] = None,
                          division: Optional[str] = None) -> str:
        """Генерирует имя файла для экспорта"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"overtime_{start_date.toString('yyyyMMdd')}_{end_date.toString('yyyyMMdd')}"

        if department and department != "Все отделы":
            filename += f"_{department}"
        if division and division != "Все подразделения":
            filename += f"_{division}"

        filename += f"_{timestamp}.xlsx"
        return filename