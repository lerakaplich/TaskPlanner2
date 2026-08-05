# services/gantt_service/gantt_export_service.py

from datetime import datetime
from typing import List, Dict, Optional, Any, Tuple

from docx.shared import Pt, Inches

from .gantt_base_service import GanttBaseService, TaskGanttData


class GanttExportService(GanttBaseService):
    """Сервис для экспорта диаграммы Ганта"""

    def __init__(self, permission_service=None, gantt_service=None):
        self.permission_service = permission_service
        self._gantt_service = gantt_service

    def can_export(self) -> bool:
        """Проверяет права на экспорт"""
        if not self.permission_service:
            return False

        from services.permissions.app_permissions import AppRole

        app_role = self.permission_service.app_manager.role

        if app_role in (AppRole.SUPER_ADMIN, AppRole.ADMIN):
            return True

        if app_role == AppRole.USER:
            if hasattr(self.permission_service, 'get_manageable_projects'):
                manageable = self.permission_service.get_manageable_projects()
                return len(manageable) > 0

        return False

    def get_all_links(self):
        """Возвращает все связи (прокси к gantt_service)"""
        if self._gantt_service:
            return self._gantt_service.get_all_links()
        return {}

    def get_task_by_id(self, task_id):
        """Возвращает задачу по ID (прокси к gantt_service)"""
        if self._gantt_service:
            return self._gantt_service.get_task_by_id(task_id)
        return None

    def export_to_image_with_period(
        self,
        canvas_widget,
        tasks: List[TaskGanttData],
        start_date: datetime,
        end_date: datetime
    ) -> Optional[str]:
        """Экспортирует диаграмму в PNG"""
        from PyQt6.QtWidgets import QFileDialog
        from PyQt6.QtGui import QPixmap
        from PyQt6.QtWidgets import QApplication

        if not canvas_widget or not tasks:
            return None

        original_tasks = canvas_widget._tasks.copy()
        original_start = canvas_widget._start_date
        original_end = canvas_widget._end_date
        original_links = canvas_widget._links.copy()

        try:
            canvas_widget.set_tasks(tasks)
            canvas_widget.set_date_range(start_date, end_date)

            # Получаем связи с типами
            from .gantt_service import GanttService
            if hasattr(canvas_widget, '_service'):
                links = canvas_widget._service.get_all_links()
                canvas_widget.set_links(links)

            canvas_widget.updateGeometry()
            canvas_widget.update()

            for _ in range(10):
                QApplication.processEvents()

            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            default_name = f"gantt_chart_{timestamp}.png"

            file_path, _ = QFileDialog.getSaveFileName(
                canvas_widget,
                "Сохранить диаграмму Ганта",
                default_name,
                "PNG Image (*.png);;JPEG Image (*.jpg *.jpeg);;All Files (*.*)"
            )

            if not file_path:
                return None

            pixmap = canvas_widget.grab()
            if pixmap.save(file_path):
                return file_path
            return None

        except Exception as e:
            print(f"❌ Ошибка экспорта PNG: {e}")
            return None

        finally:
            canvas_widget.set_tasks(original_tasks)
            canvas_widget.set_date_range(original_start, original_end)
            canvas_widget.set_links(original_links)
            canvas_widget.update()
            QApplication.processEvents()

    def _translate_priority(self, priority: str) -> str:
        """Переводит код приоритета в читаемое название"""
        priority_map = {
            "critical": "Критический",
            "high": "Высокий",
            "medium": "Средний",
            "low": "Низкий"
        }
        return priority_map.get(priority, priority)

    def export_to_excel(self, tasks: List[TaskGanttData], start_date: datetime = None, end_date: datetime = None) -> \
    Optional[str]:
        """Экспортирует диаграмму Ганта в Excel с визуальным отображением."""
        from PyQt6.QtWidgets import QFileDialog
        from datetime import datetime, timedelta
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        if not tasks:
            return None

        try:
            if start_date is None or end_date is None:
                all_dates = []
                for task in tasks:
                    all_dates.append(task.start_date)
                    all_dates.append(task.end_date)

                if all_dates:
                    start_date = min(all_dates).replace(day=1)
                    end_date = max(all_dates) + timedelta(days=14)
                else:
                    start_date = datetime.now().replace(day=1)
                    end_date = start_date + timedelta(days=30)

            current_date = start_date
            dates_list = []
            while current_date <= end_date:
                dates_list.append(current_date)
                current_date += timedelta(days=1)

            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            default_name = f"gantt_chart_{timestamp}.xlsx"

            file_path, _ = QFileDialog.getSaveFileName(
                None,
                "Экспорт диаграммы Ганта в Excel",
                default_name,
                "Excel Files (*.xlsx);;All Files (*.*)"
            )

            if not file_path:
                return None

            wb = Workbook()
            ws = wb.active
            ws.title = "Диаграмма Ганта"

            border = Border(
                left=Side(style='thin', color='E0E0E0'),
                right=Side(style='thin', color='E0E0E0'),
                top=Side(style='thin', color='E0E0E0'),
                bottom=Side(style='thin', color='E0E0E0')
            )

            header_fill = PatternFill(start_color="1B232A", end_color="1B232A", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=11)

            priority_colors = {"critical": "D22730", "high": "ccab6e", "medium": "1B232A", "low": "998664"}

            ws.merge_cells('A1:B1')
            ws['A1'] = f'Диаграмма Ганта'
            ws['A1'].font = Font(bold=True, size=14)

            ws.merge_cells('C1:E1')
            ws['C1'] = f'Период: {start_date.strftime("%d.%m.%Y")} - {end_date.strftime("%d.%m.%Y")}'
            ws['C1'].font = Font(size=10)

            ws.merge_cells('F1:H1')
            ws['F1'] = f'Задач в периоде: {len(tasks)}'
            ws['F1'].font = Font(size=10)

            headers = ['ID', 'Задача', 'Проект', 'Исполнитель', 'Приоритет', 'Прогресс']

            col_idx = 1
            for header in headers:
                cell = ws.cell(row=3, column=col_idx)
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = border
                col_idx += 1

            current_month = None
            month_start_col = col_idx

            for i, date_obj in enumerate(dates_list):
                cell = ws.cell(row=3, column=col_idx + i)
                cell.value = date_obj.day
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = border

                month_name = date_obj.strftime("%B %Y")
                if current_month != month_name:
                    if current_month is not None:
                        ws.merge_cells(
                            start_row=2, start_column=month_start_col,
                            end_row=2, end_column=col_idx + i - 1
                        )
                        month_cell = ws.cell(row=2, column=month_start_col)
                        month_cell.value = current_month
                        month_cell.fill = header_fill
                        month_cell.font = header_font
                        month_cell.alignment = Alignment(horizontal='center', vertical='center')
                        month_cell.border = border

                    current_month = month_name
                    month_start_col = col_idx + i

            if current_month:
                ws.merge_cells(
                    start_row=2, start_column=month_start_col,
                    end_row=2, end_column=col_idx + len(dates_list) - 1
                )
                month_cell = ws.cell(row=2, column=month_start_col)
                month_cell.value = current_month
                month_cell.fill = header_fill
                month_cell.font = header_font
                month_cell.alignment = Alignment(horizontal='center', vertical='center')
                month_cell.border = border

            row_num = 4

            tasks_by_project = {}
            for task in tasks:
                if task.project_name not in tasks_by_project:
                    tasks_by_project[task.project_name] = []
                tasks_by_project[task.project_name].append(task)

            for project_name, project_tasks in tasks_by_project.items():
                ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num,
                               end_column=col_idx + len(dates_list) - 1)
                project_cell = ws.cell(row=row_num, column=1)
                project_cell.value = f'📁 {project_name}'
                project_cell.font = Font(bold=True, size=12)
                project_cell.fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
                project_cell.alignment = Alignment(horizontal='left', vertical='center')
                row_num += 1

                for task in project_tasks:
                    ws.cell(row=row_num, column=1, value=task.id)
                    ws.cell(row=row_num, column=2, value=task.name)
                    ws.cell(row=row_num, column=3, value=task.project_name)
                    ws.cell(row=row_num, column=4, value=task.executor_name or "Не назначен")

                    priority_cell = ws.cell(row=row_num, column=5)
                    priority_name = self._translate_priority(task.priority)
                    priority_cell.value = priority_name
                    if task.priority in priority_colors:
                        priority_cell.fill = PatternFill(
                            start_color=priority_colors[task.priority],
                            end_color=priority_colors[task.priority],
                            fill_type="solid"
                        )
                        priority_cell.font = Font(color="FFFFFF", bold=True)

                    progress_cell = ws.cell(row=row_num, column=6)
                    progress_cell.value = f"{task.progress}%"
                    progress_cell.alignment = Alignment(horizontal='center')

                    start_offset = (task.start_date - start_date).days
                    duration = task.duration_days
                    end_offset = (task.end_date - start_date).days
                    visible_start = max(0, start_offset)
                    visible_end = min(len(dates_list) - 1, end_offset)
                    visible_duration = max(0, visible_end - visible_start + 1)

                    task_color = priority_colors.get(task.priority, "1B232A")
                    task_fill = PatternFill(start_color=task_color, end_color=task_color, fill_type="solid")

                    for i in range(visible_start, visible_start + visible_duration):
                        if 0 <= i < len(dates_list):
                            col_idx_task = 7 + i
                            cell = ws.cell(row=row_num, column=col_idx_task)
                            cell.fill = task_fill
                            cell.alignment = Alignment(horizontal='center', vertical='center')

                            if i == visible_start:
                                cell.value = task.name[:10]
                                cell.font = Font(color="FFFFFF", bold=True, size=7)
                            elif i == visible_end and visible_duration > 2:
                                if task.executor_initials:
                                    cell.value = task.executor_initials
                                    cell.font = Font(color="FFFFFF", bold=True, size=7)

                    row_num += 1

                row_num += 1

            ws.column_dimensions['A'].width = 8
            ws.column_dimensions['B'].width = 35
            ws.column_dimensions['C'].width = 25
            ws.column_dimensions['D'].width = 20
            ws.column_dimensions['E'].width = 15
            ws.column_dimensions['F'].width = 10

            for i in range(len(dates_list)):
                col_letter = get_column_letter(7 + i)
                ws.column_dimensions[col_letter].width = 4

            ws.freeze_panes = 'G4'
            ws.auto_filter.ref = f"A3:{get_column_letter(6)}{row_num - 1}"

            ws_legend = wb.create_sheet("Легенда")

            ws_legend['A1'] = 'Легенда диаграммы Ганта'
            ws_legend['A1'].font = Font(bold=True, size=14)

            ws_legend['A3'] = f'Период экспорта: {start_date.strftime("%d.%m.%Y")} - {end_date.strftime("%d.%m.%Y")}'
            ws_legend['A3'].font = Font(bold=True)

            ws_legend['A5'] = 'Приоритеты:'
            ws_legend['A5'].font = Font(bold=True)

            priority_names = {"critical": "Критический", "high": "Высокий", "medium": "Средний", "low": "Низкий"}

            row = 6
            for priority, name in priority_names.items():
                color = priority_colors.get(priority, "1B232A")
                cell = ws_legend.cell(row=row, column=1)
                cell.value = name
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                cell.font = Font(color="FFFFFF", bold=True)
                row += 1

            unique_executors = set()
            for task in tasks:
                if task.executor_name:
                    unique_executors.add(task.executor_name)

            if unique_executors:
                ws_legend.cell(row=row + 2, column=1, value="Исполнители:")
                ws_legend.cell(row=row + 2, column=1).font = Font(bold=True)

                for executor in sorted(unique_executors):
                    row += 1
                    ws_legend.cell(row=row, column=1, value=f"• {executor}")

            ws_legend.column_dimensions['A'].width = 50

            wb.save(file_path)
            print(f"✅ Excel диаграмма Ганта экспортирована: {file_path}")
            return file_path

        except ImportError as e:
            print(f"❌ Не установлена библиотека: {e}")
            from PyQt6.QtWidgets import QMessageBox
            QMessageBox.warning(
                None,
                "Ошибка экспорта",
                "Для экспорта в Excel необходима библиотека openpyxl.\n"
                "Установите её командой:\npip install openpyxl"
            )
            return None
        except Exception as e:
            print(f"❌ Ошибка экспорта в Excel: {e}")
            import traceback
            traceback.print_exc()
            return None

    def _set_cell_font(self, cell, bold=False, color=None, size=None):
        """Устанавливает шрифт для ячейки таблицы"""
        try:
            paragraph = cell.paragraphs[0]
            run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()

            if bold:
                run.bold = True
            if color:
                run.font.color.rgb = self._hex_to_rgb(color)
            if size:
                run.font.size = Pt(size)
        except Exception:
            pass

    def _auto_fit_table_columns(self, table):
        """Автоматически подбирает ширину колонок таблицы"""
        widths = [0.7, 3.0, 2.0, 2.0, 1.2, 1.5, 1.0, 1.2, 1.2, 1.0]
        for i, cell in enumerate(table.columns):
            try:
                cell.width = Inches(widths[i] if i < len(widths) else 1.0)
            except:
                pass

    def export_to_docx(self, tasks: List[TaskGanttData], start_date: datetime = None, end_date: datetime = None) -> \
    Optional[str]:
        """Экспортирует диаграмму Ганта в файл Word (DOCX)."""
        from PyQt6.QtWidgets import QFileDialog
        from datetime import datetime

        if not tasks:
            return None

        try:
            try:
                from docx import Document
                from docx.shared import Inches, Pt, RGBColor
                from docx.enum.text import WD_ALIGN_PARAGRAPH
            except ImportError:
                from PyQt6.QtWidgets import QMessageBox
                QMessageBox.warning(
                    None,
                    "Ошибка экспорта",
                    "Для экспорта в Word необходима библиотека python-docx.\n"
                    "Установите её командой:\npip install python-docx"
                )
                return None

            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            default_name = f"gantt_report_{timestamp}.docx"

            file_path, _ = QFileDialog.getSaveFileName(
                None,
                "Экспорт диаграммы Ганта в Word",
                default_name,
                "Word Documents (*.docx);;All Files (*.*)"
            )

            if not file_path:
                return None

            doc = Document()

            style = doc.styles['Normal']
            style.font.name = 'Arial'
            style.font.size = Pt(11)

            title = doc.add_heading('Диаграмма Ганта', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            doc.add_paragraph(f"Дата создания: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}")
            if start_date and end_date:
                doc.add_paragraph(
                    f"Период экспорта: {start_date.strftime('%d.%m.%Y')} - {end_date.strftime('%d.%m.%Y')}")
            doc.add_paragraph(f"Всего задач: {len(tasks)}")
            doc.add_paragraph()

            doc.add_heading('Статистика по задачам', level=1)

            priority_stats = {"critical": 0, "high": 0, "medium": 0, "low": 0}
            status_stats = {}
            executor_stats = {}

            for task in tasks:
                priority_stats[task.priority] = priority_stats.get(task.priority, 0) + 1
                status_stats[task.status] = status_stats.get(task.status, 0) + 1
                if task.executor_name:
                    executor_stats[task.executor_name] = executor_stats.get(task.executor_name, 0) + 1

            doc.add_heading('Приоритеты задач:', level=2)
            table_priority = doc.add_table(rows=len([c for c in priority_stats.values() if c > 0]) + 1, cols=2)
            table_priority.style = 'Light Grid Accent 1'

            header_cells = table_priority.rows[0].cells
            header_cells[0].text = 'Приоритет'
            header_cells[1].text = 'Количество'
            self._set_cell_font(header_cells[0], bold=True)
            self._set_cell_font(header_cells[1], bold=True)

            priority_names = {"critical": "Критический", "high": "Высокий", "medium": "Средний", "low": "Низкий"}
            row_idx = 1
            for priority, count in priority_stats.items():
                if count > 0:
                    cells = table_priority.rows[row_idx].cells
                    cells[0].text = priority_names.get(priority, priority)
                    cells[1].text = str(count)
                    row_idx += 1

            if executor_stats:
                doc.add_heading('Нагрузка на исполнителей:', level=2)
                table_executors = doc.add_table(rows=len(executor_stats) + 1, cols=2)
                table_executors.style = 'Light Grid Accent 1'

                header_cells = table_executors.rows[0].cells
                header_cells[0].text = 'Исполнитель'
                header_cells[1].text = 'Количество задач'
                self._set_cell_font(header_cells[0], bold=True)
                self._set_cell_font(header_cells[1], bold=True)

                row_idx = 1
                for executor, count in sorted(executor_stats.items()):
                    cells = table_executors.rows[row_idx].cells
                    cells[0].text = executor
                    cells[1].text = str(count)
                    row_idx += 1

            doc.add_page_break()
            doc.add_heading('Детальный список задач', level=1)

            headers = ['ID', 'Задача', 'Проект', 'Исполнитель', 'Приоритет', 'Статус', 'Прогресс', 'Дата начала',
                       'Дата окончания', 'Длительность']
            table_tasks = doc.add_table(rows=1, cols=len(headers))
            table_tasks.style = 'Light Grid Accent 1'

            header_cells = table_tasks.rows[0].cells
            for i, header in enumerate(headers):
                header_cells[i].text = header
                self._set_cell_font(header_cells[i], bold=True)
                header_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

            priority_names = {"critical": "Критический", "high": "Высокий", "medium": "Средний", "low": "Низкий"}

            for task in tasks:
                row_cells = table_tasks.add_row().cells
                row_cells[0].text = str(task.id)
                row_cells[1].text = task.name
                row_cells[2].text = task.project_name
                row_cells[3].text = task.executor_name or "Не назначен"
                row_cells[4].text = priority_names.get(task.priority, task.priority)
                row_cells[5].text = task.status
                row_cells[6].text = f"{task.progress}%"
                row_cells[7].text = task.start_date.strftime("%d.%m.%Y")
                row_cells[8].text = task.end_date.strftime("%d.%m.%Y")
                row_cells[9].text = str(task.duration_days)

                if task.priority == "critical":
                    self._set_cell_font(row_cells[4], color="D22730", bold=True)
                elif task.priority == "high":
                    self._set_cell_font(row_cells[4], color="ccab6e", bold=True)

            self._auto_fit_table_columns(table_tasks)

            # ✅ ИСПРАВЛЕНИЕ: правильно обрабатываем связи
            links = self.get_all_links()
            if links:
                task_ids = {t.id for t in tasks}
                filtered_links = {}
                for from_id, deps in links.items():
                    if from_id in task_ids:
                        filtered_to_ids = []
                        for dep in deps:
                            # Проверяем, является ли dep словарем или числом
                            if isinstance(dep, dict):
                                to_id = dep.get("successor_id")
                                link_type = dep.get("type", "FS")
                                if to_id in task_ids:
                                    filtered_to_ids.append(to_id)
                            else:
                                # Если dep - просто число
                                if dep in task_ids:
                                    filtered_to_ids.append(dep)

                        if filtered_to_ids:
                            filtered_links[from_id] = filtered_to_ids

                if filtered_links:
                    doc.add_page_break()
                    doc.add_heading('Связи между задачами', level=1)

                    link_text = ""
                    for from_id, to_ids in filtered_links.items():
                        from_task = self.get_task_by_id(from_id)
                        if from_task:
                            for to_id in to_ids:
                                to_task = self.get_task_by_id(to_id)
                                if to_task:
                                    link_text += f"• {from_task.name} → {to_task.name}\n"

                    if link_text:
                        p = doc.add_paragraph(link_text)
                        p.style.font.size = Pt(11)
                    else:
                        doc.add_paragraph("Нет связей между задачами")

            doc.save(file_path)
            print(f"✅ Word документ экспортирован: {file_path}")
            return file_path

        except Exception as e:
            print(f"❌ Ошибка экспорта в Word: {e}")
            import traceback
            traceback.print_exc()
            return None

    def get_task_info_text(self, task: TaskGanttData) -> str:
        return (
            f"Проект: {task.project_name}\n"
            f"Исполнитель: {task.executor_name or 'Не назначен'}\n"
            f"Приоритет: {task.priority}\n"
            f"Прогресс: {task.progress}%\n"
            f"Сроки: {task.start_date.strftime('%d.%m.%Y')} - {task.end_date.strftime('%d.%m.%Y')}\n"
            f"Длительность: {task.duration_days} дней"
        )